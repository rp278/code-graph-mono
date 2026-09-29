"""Builds docs/CodeGraph_Hackathon_Slides.pptx from the hackathon design doc.

Run:  python build_slides.py   (needs python-pptx)
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
LOGO = ROOT / "images" / "codegraph-logo-transparent.png"
OUT = ROOT / "docs" / "CodeGraph_Hackathon_Slides.pptx"

# ---- palette (matches the dashboard's dark theme) --------------------------
BG = RGBColor(0x0B, 0x0D, 0x14)
CARD = RGBColor(0x14, 0x18, 0x24)
CARD_EDGE = RGBColor(0x2A, 0x31, 0x45)
TEXT = RGBColor(0xF1, 0xF5, 0xF9)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
PURPLE = RGBColor(0x8B, 0x5C, 0xF6)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
GREEN = RGBColor(0x34, 0xD3, 0x99)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
PINK = RGBColor(0xF4, 0x72, 0xB6)
RED = RGBColor(0xF8, 0x71, 0x71)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333


# ---- helpers ---------------------------------------------------------------
def new_slide(notes=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(slide, x, y, w, h, content, size=16, color=TEXT, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, spacing=None):
    """content: str or list of str / (str, dict) paragraphs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, p in enumerate(paras):
        opts = {}
        if isinstance(p, tuple):
            p, opts = p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = opts.get("align", align)
        if spacing or opts.get("space"):
            para.space_after = Pt(opts.get("space", spacing))
        run = para.add_run()
        run.text = p
        run.font.name = font
        run.font.size = Pt(opts.get("size", size))
        run.font.bold = opts.get("bold", bold)
        run.font.color.rgb = opts.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=CARD, edge=CARD_EDGE, radius=0.08, shape=None,
        edge_w=1.0):
    shp = slide.shapes.add_shape(
        shape or MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if shape is None:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if edge is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = edge
        shp.line.width = Pt(edge_w)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def label_box(slide, x, y, w, h, title, sub=None, accent=CYAN, fill=CARD,
              title_size=15, sub_size=11):
    box(slide, x, y, w, h, fill=fill, edge=accent, edge_w=1.5)
    paras = [(title, {"size": title_size, "bold": True, "color": TEXT, "align": PP_ALIGN.CENTER})]
    if sub:
        paras.append((sub, {"size": sub_size, "color": MUTED, "align": PP_ALIGN.CENTER}))
    text(slide, x + 0.05, y, w - 0.1, h, paras, anchor=MSO_ANCHOR.MIDDLE)


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.75, dashed=False, arrow=True):
    c = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dashed:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    if arrow:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle")
        t.set("w", "med")
        t.set("len", "med")
    return c


def pill(slide, x, y, w, label, color=CYAN, size=11, h=0.32):
    box(slide, x, y, w, h, fill=None, edge=color, radius=0.5, edge_w=1.25)
    text(slide, x, y, w, h, label, size=size, color=color, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def header(slide, kicker, title, accent=CYAN):
    text(slide, 0.7, 0.45, 9, 0.35, kicker.upper(), size=12, color=accent, bold=True)
    text(slide, 0.7, 0.8, 11.9, 0.9, title, size=32, bold=True)
    # accent underline
    r = box(slide, 0.74, 1.62, 0.9, 0.05, fill=accent, edge=None, shape=MSO_SHAPE.RECTANGLE)


def footer(slide, n):
    text(slide, 0.7, 7.05, 6, 0.3, "CodeGraph  |  Hackathon", size=10, color=MUTED)
    text(slide, W - 1.5, 7.05, 0.8, 0.3, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, items, size=16, gap=10, color=TEXT, dot=CYAN, line_h=0.42):
    """Simple bullet list with coloured dots. Each item may wrap onto 2 lines."""
    cy = y
    for it in items:
        lines_ = it[1] if isinstance(it, tuple) else 1
        t = it[0] if isinstance(it, tuple) else it
        box(slide, x, cy + 0.11, 0.11, 0.11, fill=dot, edge=None, shape=MSO_SHAPE.OVAL)
        text(slide, x + 0.28, cy, w - 0.28, line_h * lines_, t, size=size, color=color)
        cy += line_h * lines_ + gap / 72


# ---------------------------------------------------------------------------
# 1. Title
# ---------------------------------------------------------------------------
s = new_slide(
    "One-liner: CodeGraph turns a multi-repo codebase into a live knowledge graph, then puts "
    "three workflows on top: View Graph, Ask AI and Fix Bugs. A human approves every change to code."
)
s.shapes.add_picture(str(LOGO), Inches(0.9), Inches(1.7), height=Inches(2.1))
text(s, 0.9, 3.95, 11, 1.1, "CodeGraph", size=60, bold=True)
text(s, 0.9, 5.0, 11, 0.9,
     "See, question and repair a multi-repo system from one dashboard, "
     "with a human approving every change to code.", size=22, color=MUTED)
for i, (lbl, c) in enumerate([("View Graph", CYAN), ("Ask AI", PURPLE), ("Fix Bugs", GREEN)]):
    pill(s, 0.95 + i * 1.95, 6.2, 1.75, lbl, color=c, size=13, h=0.4)
text(s, W - 4.4, 6.3, 3.7, 0.4, "Hackathon", size=14, color=MUTED, align=PP_ALIGN.RIGHT)

# ---------------------------------------------------------------------------
# 2. Problem
# ---------------------------------------------------------------------------
s = new_slide(
    "Our product spans six repos: two micro-frontends, three experience APIs (one TypeScript, two Java) "
    "and a shared design-system library. A frontend call in one repo reaches an endpoint in another, "
    "and no single tool shows that link."
)
header(s, "The problem", "Six repos, one system, no shared view", CYAN)
cards = [
    ("Links are invisible", CYAN,
     "A call in tb-discovery-mfe lands in a Java endpoint in tb-selection-xapi. "
     "No single tool shows that link across repos."),
    ("Agents guess", PURPLE,
     "General coding agents grep, miss cross-repo consumers, and change code "
     "with no audit trail."),
    ("Trust is missing", GREEN,
     "Teams won't accept unattended AI changes. They need evidence at every step, "
     "and a way to stop the agent."),
]
for i, (t, c, d) in enumerate(cards):
    x = 0.7 + i * 4.1
    box(s, x, 2.2, 3.8, 3.0, edge=c, edge_w=1.5)
    box(s, x + 0.3, 2.5, 0.5, 0.08, fill=c, edge=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.3, 2.7, 3.3, 0.5, t, size=22, bold=True)
    text(s, x + 0.3, 3.35, 3.25, 1.8, d, size=15, color=MUTED)

# repo strip
text(s, 0.7, 5.55, 6, 0.3, "THE SIX REPOS", size=11, color=MUTED, bold=True)
repos = [("tb-common-mfe", BLUE), ("tb-discovery-mfe", BLUE), ("tb-marketing-xapi", AMBER),
         ("tb-discovery-xapi", PINK), ("tb-selection-xapi", PINK), ("kairos-fabric", GREEN)]
for i, (r, c) in enumerate(repos):
    pill(s, 0.7 + i * 2.03, 5.95, 1.93, r, color=c, size=11, h=0.42)
text(s, 0.7, 6.5, 12, 0.3, "Blue: micro-frontends   Amber: TypeScript API   Pink: Java APIs   Green: design-system library",
     size=11, color=MUTED)
footer(s, 2)

# ---------------------------------------------------------------------------
# 3. Solution
# ---------------------------------------------------------------------------
s = new_slide(
    "Three workflows on one foundation: a cross-repo knowledge graph built by Graphify in Neo4j, a React "
    "dashboard drawing it with React Flow, and Cursor SDK agents grounded by the graph."
)
header(s, "The solution", "One graph, three workflows", PURPLE)
wf = [
    ("View Graph", CYAN, "See",
     "Explore how repos, files, components, endpoints, tables and packages connect, including links across repos.",
     "No code changes"),
    ("Ask AI", PURPLE, "Question",
     "Ask in plain English and get an answer with real file names and line numbers.",
     "Read-only"),
    ("Fix Bugs", GREEN, "Repair",
     "Paste an error. An agent finds the cause, proves it with a failing test, fixes it minimally and opens a PR "
     "through 4 human gates.",
     "Branch only. Never merges."),
]
for i, (t, c, verb, d, badge) in enumerate(wf):
    x = 0.7 + i * 4.1
    box(s, x, 2.0, 3.8, 3.55, edge=c, edge_w=1.5)
    text(s, x + 0.3, 2.2, 3.2, 0.3, verb.upper(), size=12, color=c, bold=True)
    text(s, x + 0.3, 2.5, 3.3, 0.6, t, size=26, bold=True)
    text(s, x + 0.3, 3.2, 3.25, 1.6, d, size=14, color=MUTED)
    pill(s, x + 0.3, 4.95, 2.6, badge, color=c, size=11, h=0.34)
# foundation bar
box(s, 0.7, 5.85, 11.9, 0.95, fill=CARD, edge=CARD_EDGE)
text(s, 0.95, 5.9, 3, 0.3, "SHARED FOUNDATION", size=11, color=MUTED, bold=True)
found = [("Cross-repo knowledge graph", "Graphify + Neo4j", CYAN),
         ("React dashboard", "React Flow", BLUE),
         ("Cursor SDK agents", "grounded by the graph", PURPLE)]
for i, (a, b, c) in enumerate(found):
    x = 0.95 + i * 3.95
    box(s, x, 6.23, 0.09, 0.42, fill=c, edge=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.2, 6.2, 3.6, 0.5, [(a, {"size": 14, "bold": True}), (b, {"size": 11, "color": MUTED})])
footer(s, 3)

# ---------------------------------------------------------------------------
# 4. Architecture
# ---------------------------------------------------------------------------
s = new_slide(
    "Graphify extracts nodes and edges from the six repos and pushes them to Neo4j. FastAPI serves the "
    "graph to the React dashboard and hosts the two agents. Ask AI is read-only; Fix Bugs edits, tests "
    "and opens PRs on GitHub, and records progress in state.json."
)
header(s, "Architecture", "How the pieces fit together", CYAN)
# top row
label_box(s, 0.7, 2.25, 1.9, 1.0, "6 work repos", "source code", BLUE)
label_box(s, 3.15, 2.25, 1.9, 1.0, "Graphify", "extract and link", CYAN)
label_box(s, 5.6, 2.25, 1.9, 1.0, "Neo4j", "knowledge graph", GREEN)
label_box(s, 8.05, 2.25, 1.9, 1.0, "FastAPI", "port 8000", PURPLE)
label_box(s, 10.5, 2.25, 2.1, 1.0, "React dashboard", "port 5173", PINK)
for x in (2.6, 5.05, 7.5):
    line(s, x, 2.75, x + 0.55, 2.75)
line(s, 9.95, 2.62, 10.5, 2.62)
line(s, 10.5, 2.88, 9.95, 2.88)
# agents
label_box(s, 6.2, 3.95, 2.3, 0.9, "Ask AI agent", "read-only", PURPLE)
label_box(s, 9.0, 4.4, 2.3, 1.0, "Fix Bugs agent", "gated", GREEN)
line(s, 8.6, 3.25, 7.35, 3.95)
line(s, 9.4, 3.25, 10.15, 4.4)
# outputs
label_box(s, 0.7, 4.4, 2.3, 1.0, "Repo checkouts", "read and edit", BLUE)
label_box(s, 9.0, 6.0, 1.9, 0.75, "GitHub", "open PR", AMBER, title_size=14)
label_box(s, 11.2, 6.0, 1.4, 0.75, "state.json", None, AMBER, title_size=13)
line(s, 6.2, 4.4, 3.0, 4.6, color=PURPLE)
line(s, 9.0, 5.2, 3.0, 5.2, color=GREEN, dashed=True)
line(s, 9.95, 5.4, 9.95, 6.0)
line(s, 10.75, 5.4, 11.7, 6.0)
text(s, 3.2, 4.05, 2.9, 0.3, "read only", size=11, color=PURPLE, align=PP_ALIGN.CENTER)
text(s, 3.2, 5.3, 5.5, 0.3, "edit, test, git", size=11, color=GREEN, align=PP_ALIGN.CENTER)
# stack
text(s, 0.7, 6.05, 8, 0.3, "STACK", size=11, color=MUTED, bold=True)
text(s, 0.7, 6.35, 8, 0.7,
     "Python + FastAPI, Neo4j, tree-sitter, React 19 + Vite, @xyflow/react + dagre, Cursor SDK, gh CLI",
     size=12, color=MUTED)
footer(s, 4)

# ---------------------------------------------------------------------------
# 5. Graphify
# ---------------------------------------------------------------------------
s = new_slide(
    "Graphify is a customized fork of the open-source safishamsi/graphify with a deterministic multi-repo path. "
    "No LLM at build time: same code, same graph. LLMs only read the graph. Run: "
    "python -m graphify.codegraph --repos repos.json --push, or press Rebuild in the dashboard. "
    "Stats are a snapshot: about 13,200 nodes and 35,900 edges."
)
header(s, "Graphify", "A deterministic graph builder: no LLM at build time", CYAN)
steps = [
    ("1", "Collect files", "skip node_modules, dist, build…"),
    ("2", "Base AST", "tree-sitter: functions, classes, calls"),
    ("3", "Framework rules", "Fastify, React, Next.js, Spring, SQL"),
    ("4", "Tag with repo", "ids prefixed <repo>::"),
]
steps2 = [
    ("5", "Package graph", "package.json to package nodes"),
    ("6", "Link pass", "resolve edges across repos"),
    ("7", "Dedupe, push", "validate, MERGE into Neo4j"),
]
text(s, 0.7, 1.95, 4, 0.3, "PER REPO", size=11, color=MUTED, bold=True)
for i, (n, t, d) in enumerate(steps):
    x = 0.7 + i * 3.05
    box(s, x, 2.3, 2.8, 1.25, edge=CYAN)
    text(s, x + 0.15, 2.35, 0.5, 0.4, n, size=20, bold=True, color=CYAN)
    text(s, x + 0.6, 2.4, 2.1, 0.35, t, size=14, bold=True)
    text(s, x + 0.2, 2.85, 2.45, 0.7, d, size=11, color=MUTED)
    if i < 3:
        line(s, x + 2.8, 2.92, x + 3.05, 2.92)
text(s, 0.7, 3.8, 4, 0.3, "ALL REPOS TOGETHER", size=11, color=MUTED, bold=True)
for i, (n, t, d) in enumerate(steps2):
    x = 0.7 + i * 3.05
    box(s, x, 4.15, 2.8, 1.25, edge=PURPLE)
    text(s, x + 0.15, 4.2, 0.5, 0.4, n, size=20, bold=True, color=PURPLE)
    text(s, x + 0.6, 4.25, 2.1, 0.35, t, size=14, bold=True)
    text(s, x + 0.2, 4.7, 2.45, 0.7, d, size=11, color=MUTED)
    if i < 2:
        line(s, x + 2.8, 4.77, x + 3.05, 4.77)
# stats panel
box(s, 9.95, 3.85, 2.65, 1.55, fill=CARD, edge=CARD_EDGE)
text(s, 10.1, 3.9, 2.4, 0.3, "CURRENT GRAPH", size=10, color=MUTED, bold=True)
text(s, 10.1, 4.2, 2.4, 0.6, "13.2k nodes", size=22, bold=True, color=CYAN)
text(s, 10.1, 4.75, 2.4, 0.6, "35.9k edges", size=22, bold=True, color=PURPLE)
# bottom strip
box(s, 0.7, 5.75, 11.9, 1.0, fill=CARD, edge=CARD_EDGE)
text(s, 0.95, 5.8, 11.4, 0.3, "WHY DETERMINISTIC", size=11, color=MUTED, bold=True)
text(s, 0.95, 6.1, 11.4, 0.6,
     "The same code always produces the same graph. LLMs only read the graph (Ask AI, Fix Bugs). "
     "They never create it. Rebuilds are idempotent per repo.", size=14)
footer(s, 5)

# ---------------------------------------------------------------------------
# 6. Link pass
# ---------------------------------------------------------------------------
s = new_slide(
    "Extractors observe one file at a time; linkers connect. An extractor can only say component X calls "
    "GET /api/products, target unknown, so it records a pending edge. After all repos are extracted the link "
    "pass builds a global (method, path) registry and resolves each pending call. Also follows API-client "
    "wrappers, service calls via env-var base URLs, and Next.js proxy routes. Resolved cross-repo edges are "
    "INFERRED (dashed); calls_service is EXTRACTED."
)
header(s, "The key idea", "Extractors observe. Linkers connect.", PURPLE)
# diagram
box(s, 0.7, 2.1, 3.7, 1.8, edge=BLUE, edge_w=1.5)
text(s, 0.9, 2.15, 3.3, 0.3, "tb-discovery-mfe", size=12, color=BLUE, bold=True)
label_box(s, 1.0, 2.6, 3.1, 0.9, "ProductList", "component  |  fetch('/api/products')", BLUE, title_size=14, sub_size=10)
box(s, 8.9, 2.1, 3.7, 1.8, edge=PINK, edge_w=1.5)
text(s, 9.1, 2.15, 3.3, 0.3, "tb-selection-xapi", size=12, color=PINK, bold=True)
label_box(s, 9.2, 2.6, 3.1, 0.9, "GET /api/products", "endpoint  |  Spring controller", PINK, title_size=14, sub_size=10)
line(s, 4.1, 3.05, 9.2, 3.05, color=PURPLE, width=2.5, dashed=True)
pill(s, 5.55, 2.6, 2.2, "fetches  |  INFERRED", color=PURPLE, size=11, h=0.34)
text(s, 4.5, 3.2, 4.4, 0.3, "resolved by the global link pass", size=11, color=MUTED, align=PP_ALIGN.CENTER)
# three ways
text(s, 0.7, 4.2, 6, 0.3, "THREE WAYS THE LINK PASS CONNECTS REPOS", size=11, color=MUTED, bold=True)
ways = [
    ("Pending calls", CYAN,
     "Match (method, path) to a global endpoint registry. Tolerates :param segments and API-client wrappers."),
    ("Service calls", PURPLE,
     "An env var like SELECTION_XAPI_BASE_URL becomes a calls_service edge to a service node, "
     "then narrows to real endpoints."),
    ("Proxy routes", GREEN,
     "A Next.js catch-all route links to the repo it forwards to (proxies_to). A bare /api/:path* is never linked."),
]
for i, (t, c, d) in enumerate(ways):
    x = 0.7 + i * 4.1
    box(s, x, 4.55, 3.8, 2.2, edge=c)
    text(s, x + 0.25, 4.7, 3.3, 0.4, t, size=18, bold=True, color=c)
    text(s, x + 0.25, 5.2, 3.35, 1.5, d, size=13, color=MUTED)
footer(s, 6)

# ---------------------------------------------------------------------------
# 7. View Graph
# ---------------------------------------------------------------------------
s = new_slide(
    "13,000 nodes on one canvas is unreadable, so we never draw the whole graph. Explore starts on the "
    "All-repos map: one box per repo and one merged arrow per repo pair, with counts. Pick a node and you see "
    "its neighborhood: the node plus everything one hop away, in any repo. Click a neighbor to walk. The "
    "Tree view browses a repo like a file explorer. React Flow gives pan/zoom/minimap; dagre lays out."
)
header(s, "Workflow 1", "View Graph: progressive disclosure on React Flow", CYAN)
# left: the drill-down ladder
levels = [
    ("All-repos map", "One box per repo, one arrow per repo pair, labelled with counts", CYAN),
    ("Pick a repo", "Repo in the middle, entry points around it: packages, pages, endpoints", BLUE),
    ("Neighborhood", "A node plus everything one hop away, in any repo. Click to walk.", PURPLE),
]
for i, (t, d, c) in enumerate(levels):
    y = 2.05 + i * 1.45
    box(s, 0.7, y, 6.1, 1.25, edge=c, edge_w=1.5)
    text(s, 0.95, y + 0.08, 0.6, 0.6, str(i + 1), size=28, bold=True, color=c)
    text(s, 1.6, y + 0.15, 5.0, 0.4, t, size=18, bold=True)
    text(s, 1.6, y + 0.58, 5.0, 0.6, d, size=12, color=MUTED)
    if i < 2:
        line(s, 3.75, y + 1.25, 3.75, y + 1.45, color=MUTED)
# right: neighborhood mini diagram
box(s, 7.2, 2.05, 5.4, 4.05, fill=CARD, edge=CARD_EDGE)
text(s, 7.4, 2.12, 5, 0.3, "NEIGHBORHOOD OF ProductList (ILLUSTRATIVE)", size=10, color=MUTED, bold=True)
label_box(s, 9.0, 3.65, 1.8, 0.75, "ProductList", "focus", PURPLE, fill=RGBColor(0x24, 0x1B, 0x4B),
          title_size=14, sub_size=10)
label_box(s, 7.4, 2.7, 1.5, 0.6, "ProductsPage", None, BLUE, title_size=11)
label_box(s, 7.4, 4.7, 1.5, 0.6, "ProductList.tsx", None, MUTED, title_size=11)
label_box(s, 11.0, 2.7, 1.45, 0.6, "formatPrice()", None, GREEN, title_size=11)
label_box(s, 10.95, 4.7, 1.55, 0.6, "GET /api/products", None, PINK, title_size=10)
line(s, 8.9, 3.05, 9.3, 3.65, color=BLUE)
line(s, 8.9, 5.0, 9.4, 4.4, color=MUTED)
line(s, 10.6, 3.65, 11.2, 3.3, color=GREEN)
line(s, 10.6, 4.4, 11.2, 4.75, color=PINK, dashed=True)
text(s, 7.4, 5.45, 5.0, 0.6, "Left: used by / parent.   Right: uses / children.\nDashed = inferred, other repo.",
     size=11, color=MUTED)
# bottom stats
stats = [("13.2k", "nodes"), ("35.9k", "edges"), ("6", "repos"), ("88", "endpoints"), ("839", "components")]
for i, (n, l) in enumerate(stats):
    x = 0.7 + i * 2.4
    text(s, x, 6.25, 2.3, 0.5, n, size=24, bold=True, color=CYAN)
    text(s, x, 6.7, 2.3, 0.3, l, size=11, color=MUTED)
footer(s, 7)

# ---------------------------------------------------------------------------
# 8. Ask AI
# ---------------------------------------------------------------------------
s = new_slide(
    "Ask AI pre-fetches a small keyword-based subgraph (up to 8 keywords, 25 nodes, 6 neighbors each) and puts "
    "it in the prompt so the agent knows where to look. The agent's toolset is restricted to read, grep, glob, "
    "ls; that is enforced by the SDK at agent creation, not requested in the prompt. If the graph and a file "
    "disagree, the agent trusts the file and says so. Runs in a background thread; UI polls live status."
)
header(s, "Workflow 2", "Ask AI: answers with real files and line numbers", PURPLE)
flow = [("Question", "plain English", BLUE), ("Graph context", "keyword subgraph", CYAN),
        ("Agent", "opens real files", PURPLE), ("Answer", "file:line + chips", GREEN)]
for i, (t, d, c) in enumerate(flow):
    x = 0.7 + i * 3.1
    label_box(s, x, 2.15, 2.6, 1.05, t, d, c)
    if i < 3:
        line(s, x + 2.6, 2.67, x + 3.1, 2.67)
# read-only toolset
box(s, 0.7, 3.65, 5.8, 3.0, edge=PURPLE, edge_w=1.5)
text(s, 0.95, 3.75, 5.3, 0.4, "Read-only, enforced by the SDK", size=18, bold=True)
for i, tname in enumerate(["read", "grep", "glob", "ls"]):
    pill(s, 0.95 + i * 1.35, 4.35, 1.2, tname, color=GREEN, size=13, h=0.4)
text(s, 0.95, 4.95, 5.3, 1.6,
     "The agent has no shell, edit or delete tool. Read-only is set at agent creation, not requested in the prompt.",
     size=13, color=MUTED)
# design notes
bullets(s, 6.9, 3.75, 5.7, [
    ("Graph is a map: if it disagrees with a file, the agent trusts the file", 2),
    ("One conversation is one agent: follow-ups keep context, and resume after restart", 2),
    ("Live status in the UI: \"Running grep...\" instead of a bare spinner", 2),
    ("No graph match? The agent searches the code itself and says so", 2),
], size=14, gap=8, dot=PURPLE, line_h=0.32)
footer(s, 8)

# ---------------------------------------------------------------------------
# 9. Fix Bugs pipeline
# ---------------------------------------------------------------------------
s = new_slide(
    "Paste an error; the agent works through four stages with four mandatory human gates. Stages 1 to 3 may "
    "not touch source: only test files (documents go to codegraph/pipeline/). Stage 3 exists to falsify the "
    "hypothesis; if the failing test fails for a different reason, we return to root cause. When Gate 4 is "
    "approved the run is done: the pipeline stops at an open PR and never merges. A person reviews and merges."
)
header(s, "Workflow 3", "Fix Bugs: test-proven fixes, four human gates", GREEN)
stages = [
    ("1", "Analyze\nand locate", "bugfix.md", BLUE),
    ("2", "Root\ncause", "rootcause.md", PURPLE),
    ("3", "Reproduce\non unfixed code", "repro.md", AMBER),
    ("4", "Fix and\nopen PR", "fix.md", GREEN),
]
bw, gap = 2.65, 0.43
for i, (n, t, art, c) in enumerate(stages):
    x = 0.7 + i * (bw + gap)
    box(s, x, 2.15, bw, 1.75, edge=c, edge_w=1.5)
    text(s, x + 0.15, 2.2, 0.5, 0.4, n, size=20, bold=True, color=c)
    text(s, x + 0.15, 2.65, bw - 0.3, 0.75, t, size=15, bold=True)
    text(s, x + 0.15, 3.45, bw - 0.3, 0.3, art, size=10, color=MUTED)
    if i < 3:
        gx = x + bw
        line(s, gx, 3.0, gx + gap, 3.0, color=RED, width=2.25)
        pill(s, gx - 0.05, 2.55, gap + 0.1, f"G{i + 1}", color=RED, size=10, h=0.3)
# Gate 4 sits between the last stage and the finished run
last_x = 0.7 + 3 * (bw + gap) + bw / 2
line(s, last_x, 3.9, last_x, 4.6, color=RED, width=2.25)
pill(s, last_x + 0.1, 4.05, 0.55, "G4", color=RED, size=10, h=0.3)
text(s, 0.7, 4.05, 9.6, 0.3,
     "Each gate is a mandatory human approval. Stage 3 falls back to Stage 2 if the hypothesis was wrong.",
     size=11, color=MUTED)
# done
box(s, 0.7, 4.6, 11.9, 0.9, fill=CARD, edge=GREEN, edge_w=1.5)
text(s, 0.95, 4.65, 5, 0.3, "RUN DONE: PR LEFT OPEN", size=11, color=GREEN, bold=True)
text(s, 0.95, 4.95, 11.4, 0.5,
     "After Gate 4, merge_status: not_merged. A person reads the diff and the artifacts, then merges (or doesn't).",
     size=14)
# rules
rules = [("No source edits before Gate 3", "Only test files in the repo"),
         ("Scope fence", "Fixed at Gate 2; widening re-opens it"),
         ("Never merges", "Agent may not run gh pr merge")]
for i, (t, d) in enumerate(rules):
    x = 0.7 + i * 4.1
    box(s, x, 5.75, 3.8, 1.05, edge=CARD_EDGE)
    text(s, x + 0.2, 5.82, 3.4, 0.4, t, size=15, bold=True)
    text(s, x + 0.2, 6.25, 3.4, 0.5, d, size=12, color=MUTED)
footer(s, 9)

# ---------------------------------------------------------------------------
# 10. Different from "ask an agent"
# ---------------------------------------------------------------------------
s = new_slide(
    "Why not just ask an agent to fix it? The hypothesis is a claim and stage 3 exists to falsify it. "
    "Preservation tests come from the graph's blast radius, including consumers in other repos. "
    "Headless approvals: the agent writes a pending_gate marker in state.json and stops; the dashboard "
    "shows it; approve or revise resumes the same agent. Run controls in the dashboard: Pause cancels the "
    "in-flight step (the agent re-checks its work on Resume), Restart starts the same run over with all gates "
    "cleared (repos are not cleaned up), Delete removes only the run record."
)
header(s, "What makes it different", "Not just \"ask an agent to fix it\"", GREEN)
pts = [
    ("Test-first and falsifiable", "The root cause is a claim. A test that fails for the wrong reason sends the run back to Stage 2.", GREEN),
    ("Preservation from the graph", "Blast radius, including consumers in other repos, decides what must not change.", CYAN),
    ("Every gate is mandatory", "Silence or a blanket \"looks fine\" is not approval. Rejection loops back.", RED),
    ("Scope fence", "Gate 2 fixes which files and functions may change. Leaving it means re-approval.", AMBER),
    ("\"No test\" is stated, never silent", "For races or env-only bugs, repro.md gives the reason and manual steps.", PURPLE),
    ("You stay in control of a run", "Pause and resume, restart from scratch, or delete it. state.json survives API restarts.", BLUE),
]
for i, (t, d, c) in enumerate(pts):
    col, row = i % 3, i // 3
    x, y = 0.7 + col * 4.1, 2.0 + row * 2.35
    box(s, x, y, 3.8, 2.1, edge=c)
    box(s, x, y + 0.25, 0.09, 0.5, fill=c, edge=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.3, y + 0.2, 3.35, 0.7, t, size=17, bold=True)
    text(s, x + 0.3, y + 0.95, 3.35, 1.1, d, size=12, color=MUTED)
footer(s, 10)

# ---------------------------------------------------------------------------
# 11. Safety
# ---------------------------------------------------------------------------
s = new_slide(
    "Ask AI's read-only guarantee is technical: the toolset has no edit tool. Fix Bugs is different: its "
    "guardrails are rules the agent must follow, because it has a shell and git. That's why gates, scope "
    "fence and no-merge are backed by a human at every step and the last decision stays with a person."
)
header(s, "Safety model", "A human keeps the last word on every change", RED)
rows = [
    ("View Graph", "No source edits", "Read endpoints. Rebuild only rewrites the graph (fast-forward git pull).", CYAN),
    ("Ask AI", "No", "SDK-enforced toolset: read, grep, glob, ls. No shell or edit tool exists.", PURPLE),
    ("Fix Bugs", "Branch only", "4 gates, no edits before Gate 3, scope fence, minimal change, no merge.", GREEN),
]
text(s, 0.7, 1.95, 3, 0.3, "WORKFLOW", size=11, color=MUTED, bold=True)
text(s, 3.2, 1.95, 3, 0.3, "CAN IT CHANGE CODE?", size=11, color=MUTED, bold=True)
text(s, 6.0, 1.95, 3, 0.3, "GUARDRAIL", size=11, color=MUTED, bold=True)
for i, (wfn, can, g, c) in enumerate(rows):
    y = 2.3 + i * 1.15
    box(s, 0.7, y, 11.9, 1.0, edge=c)
    text(s, 0.95, y, 2.2, 1.0, wfn, size=18, bold=True, color=c, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.2, y, 2.6, 1.0, can, size=15, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 6.0, y, 6.4, 1.0, g, size=13, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.7, 5.9, 11.9, 0.3, "OTHER CONTROLS", size=11, color=MUTED, bold=True)
for i, t in enumerate(["Optional bearer token on every route except /health",
                       "CORS origin allowlist",
                       "Secrets only in .env files",
                       "LLMs never create graph nodes or edges"]):
    pill(s, 0.7 + (i % 2) * 6.05, 6.25 + (i // 2) * 0.42, 5.85, t, color=MUTED, size=11, h=0.34)
footer(s, 11)

# ---------------------------------------------------------------------------
# 12. Future + close
# ---------------------------------------------------------------------------
s = new_slide("Close: see, question and repair the system, with evidence at each step and a person deciding what merges.")
header(s, "What's next", "Future work", PURPLE)
future = [
    ("Incremental rebuilds", "Instead of full rebuilds, plus repo-relative paths", CYAN),
    ("More extractors", "Other frameworks and message queues, plus unresolved-fetch checks", BLUE),
    ("Graph diffs and trails", "Saved trails, shareable node links, diff between rebuilds", PURPLE),
    ("Semantic search", "Better graph context for Ask AI", PINK),
    ("Multi-user approvals", "Identity and per-repo owners", AMBER),
    ("Isolated worktrees + metrics", "Safer runs, and time per gate, revise rate, falsified hypotheses", GREEN),
]
for i, (t, d, c) in enumerate(future):
    col, row = i % 3, i // 3
    x, y = 0.7 + col * 4.1, 2.0 + row * 1.5
    box(s, x, y, 3.8, 1.3, edge=c)
    text(s, x + 0.25, y + 0.12, 3.3, 0.4, t, size=16, bold=True, color=c)
    text(s, x + 0.25, y + 0.55, 3.3, 0.7, d, size=12, color=MUTED)
# closing band
box(s, 0.7, 5.15, 11.9, 1.6, fill=CARD, edge=CARD_EDGE)
s.shapes.add_picture(str(LOGO), Inches(0.95), Inches(5.3), height=Inches(1.3))
text(s, 3.2, 5.3, 9.2, 0.6, "See it. Question it. Repair it.", size=28, bold=True)
text(s, 3.2, 5.95, 9.2, 0.7,
     "A live graph of the system, answers with evidence, and fixes a human approves at every gate.",
     size=15, color=MUTED)
footer(s, 12)

prs.save(OUT)
print("saved", OUT)
